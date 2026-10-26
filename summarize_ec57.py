#!/usr/bin/env python3
"""Summarise an EC57 result folder (epicmp rhythm reports + bxb beat reports) into tables.

    python3 summarize_ec57.py                        # ec57_results/ next to this file
    python3 summarize_ec57.py ec57_results           # one folder
    python3 summarize_ec57.py run_a/ run_b/          # several folders side by side
    python3 summarize_ec57.py --records              # also per-record CSV + worst records
    python3 summarize_ec57.py --out /tmp/sum         # where the CSV / XLSX / MD go

Input layout (what evaluate.py / `python3 -m ecgr.rhythm ec57` write):
    <dir>/<db>/<db>_<CLASS>_report_line.out          epicmp  (rhythm episodes)
    <dir>/beats/<db>/<db>_QRS_report_line.out        bxb     (beat labels)

Output (in --out, default = the first result folder):
    ec57_summary.csv          one row per database x class (episode + duration Se/+P/F1, counts)
    ec57_beats_summary.csv    one row per database (Q/V/S Se/+P, totals)
    ec57_summary.xlsx         both tables (+ per-record sheets with --records); needs openpyxl
    ec57_summary.md           the printed tables as markdown
    ec57_records_*.csv        per-record rows (--records)

Only the standard library is required; openpyxl is optional (XLSX is skipped without it).
'-' is kept where a database has no reference of that class (not turned into 0).
"""
import argparse
import csv
import glob
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPORT_SUFFIX = '_report_line.out'
CLASS_ORDER = ['AFIB', 'SVT', 'VT', 'AVB']
DB_ORDER = ['mitdb', 'afdb', 'escdb', 'ahadb', 'nstdb', 'rhythm_eval']

# The reference product's table ("rhythm 3.0.6"), (Se, PPV) - same as ecgr/rhythm/ec57.py.
TARGET_TABLE = {
    ('mitdb', 'AFIB'): {'D': (98, 90), 'E': (80, 92)},
    ('afdb', 'AFIB'): {'D': (96, 99), 'E': (81, 93)},
    ('mitdb', 'SVT'): {'D': (61, 24), 'E': (62, 19)},
    ('mitdb', 'VT'): {'D': (75, 57), 'E': (82, 58)},
    ('mitdb', 'AVB'): {'D': (98, 87), 'E': (100, 100)},
}

RHYTHM_COLUMNS = ['db', 'class', 'records', 'E_Se', 'E_+P', 'E_F1', 'D_Se', 'D_+P', 'D_F1',
                  'TPs', 'FN', 'TPp', 'FP', 'ref_duration_s', 'test_duration_s']
BEAT_METRICS = ['Q_Se', 'Q_+P', 'V_Se', 'V_+P', 'S_Se', 'S_+P']
BEAT_COLUMNS = ['db', 'records'] + BEAT_METRICS + ['V_F1', 'S_F1', 'total_QRS', 'total_VEB',
                                                   'total_SVEB', 'RR_err']
BXB_COUNTS = ["Nn'", "Sn'", "Vn'", "Fn'", "On'", 'Ns', 'Ss', 'Vs', "Fs'", "Os'", 'Nv', 'Sv',
              'Vv', "Fv'", "Ov'", "No'", "So'", "Vo'", "Fo'"]


# ----------------------------------------------------------------------------- helpers
def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def f1(se, pp):
    se, pp = num(se), num(pp)
    if se is None or pp is None or se + pp == 0:
        return '-'
    return f"{2 * se * pp / (se + pp):.1f}"


def duration_s(tok):
    """'1:43:46.780' / '24:29.556' / '0.000' -> seconds."""
    parts = tok.split(':')
    try:
        s = 0.0
        for p in parts:
            s = s * 60 + float(p)
        return round(s, 3)
    except ValueError:
        return None


def order(items, preferred):
    rank = {k: i for i, k in enumerate(preferred)}
    return sorted(set(items), key=lambda k: (rank.get(k, len(rank)), k))


# ----------------------------------------------------------------------------- epicmp
def parse_epicmp(path):
    """Gross / Sum / Summary lines + per-record rows of one epicmp report."""
    gross, records = {}, []
    with open(path) as f:
        for line in f:
            t = line.split()
            if not t or line.startswith('(') or line.startswith('Record') or line.startswith('_'):
                continue
            if t[0] == 'Gross' and len(t) >= 5:
                gross.update(zip(['E_Se', 'E_+P', 'D_Se', 'D_+P'], t[1:5]))
            elif t[0] == 'Sum' and len(t) >= 7:
                gross.update(zip(['TPs', 'FN', 'TPp', 'FP'], t[1:5]))
                gross['ref_duration_s'] = duration_s(t[5])
                gross['test_duration_s'] = duration_s(t[6])
            elif line.startswith('Summary of results from'):
                gross['records'] = t[4]
            elif t[0] != 'Average' and len(t) >= 11:
                records.append({
                    'record': t[0], 'TPs': t[1], 'FN': t[2], 'TPp': t[3], 'FP': t[4],
                    'E_Se': t[5], 'E_+P': t[6], 'D_Se': t[7], 'D_+P': t[8],
                    'ref_duration_s': duration_s(t[9]), 'test_duration_s': duration_s(t[10])})
    return gross, records


def rhythm_rows(ec57_dir):
    rows, per_record = [], []
    for path in sorted(glob.glob(os.path.join(ec57_dir, '*', '*' + REPORT_SUFFIX))):
        db = os.path.basename(os.path.dirname(path))
        if db == 'beats':
            continue
        name = os.path.basename(path)[:-len(REPORT_SUFFIX)]
        cls = name[len(db) + 1:] if name.startswith(db + '_') else name
        gross, recs = parse_epicmp(path)
        row = {'db': db, 'class': cls, **gross}
        row['E_F1'] = f1(row.get('E_Se'), row.get('E_+P'))
        row['D_F1'] = f1(row.get('D_Se'), row.get('D_+P'))
        rows.append(row)
        for r in recs:
            per_record.append({'db': db, 'class': cls, **r,
                               'E_F1': f1(r['E_Se'], r['E_+P']), 'D_F1': f1(r['D_Se'], r['D_+P'])})
    return rows, per_record


# ----------------------------------------------------------------------------- bxb
def parse_bxb(path):
    gross, records = {}, []
    with open(path) as f:
        for line in f:
            t = line.split()
            if not t or line.startswith('Record') or line.startswith('_'):
                continue
            if t[0] == 'Gross' and len(t) >= 7:
                gross.update(zip(BEAT_METRICS, t[1:7]))
            elif t[0] == 'Average' and len(t) >= 8:
                gross['RR_err'] = t[7]
            elif line.startswith('Total QRS complexes:'):
                gross['total_QRS'], gross['total_VEB'], gross['total_SVEB'] = t[3], t[6], t[9]
            elif line.startswith('Summary of results from'):
                gross['records'] = t[4]
            elif t[0] not in ('Sum', 'Average') and len(t) >= 1 + len(BXB_COUNTS) + 6:
                r = {'record': t[0]}
                r.update(zip(BXB_COUNTS, t[1:1 + len(BXB_COUNTS)]))
                r.update(zip(BEAT_METRICS, t[1 + len(BXB_COUNTS):7 + len(BXB_COUNTS)]))
                r['RR_err'] = t[7 + len(BXB_COUNTS)] if len(t) > 7 + len(BXB_COUNTS) else '-'
                records.append(r)
    return gross, records


def beat_rows(ec57_dir):
    rows, per_record = [], []
    for path in sorted(glob.glob(os.path.join(ec57_dir, 'beats', '*', '*_QRS' + REPORT_SUFFIX))):
        db = os.path.basename(os.path.dirname(path))
        gross, recs = parse_bxb(path)
        row = {'db': db, **gross}
        row['V_F1'] = f1(row.get('V_Se'), row.get('V_+P'))
        row['S_F1'] = f1(row.get('S_Se'), row.get('S_+P'))
        rows.append(row)
        for r in recs:
            per_record.append({'db': db, **r, 'V_F1': f1(r['V_Se'], r['V_+P']),
                               'S_F1': f1(r['S_Se'], r['S_+P'])})
    return rows, per_record


# ----------------------------------------------------------------------------- tables
def mark(row, key, db, cls):
    """'*' when both Se and +P reach the product target, '!' when not, '' without a target."""
    t = (TARGET_TABLE.get((db, cls)) or {}).get(key)
    se, pp = num(row.get(f'{key}_Se')), num(row.get(f'{key}_+P'))
    if not t or se is None or pp is None:
        return ''
    return '*' if se >= t[0] and pp >= t[1] else '!'


def rhythm_table(rows, tag=None):
    """Markdown table: class x (Duration | Episode), one Se / +P / F1 triple per database."""
    by = {(r['db'], r['class']): r for r in rows}
    dbs = order([r['db'] for r in rows], DB_ORDER)
    classes = order([r['class'] for r in rows], CLASS_ORDER)
    head = ['class', 'kind'] + [f"{db} Se/+P/F1" for db in dbs]
    lines = []
    if tag:
        lines.append(f"**{tag}**\n")
    lines += ['| ' + ' | '.join(head) + ' |', '|' + '---|' * len(head)]
    for cls in classes:
        for kind, key in (('Duration', 'D'), ('Episode', 'E')):
            cells = []
            for db in dbs:
                r = by.get((db, cls))
                if r is None:
                    cells.append('')
                    continue
                cells.append(f"{r[f'{key}_Se']} / {r[f'{key}_+P']} / {r[f'{key}_F1']}"
                             f"{(' ' + mark(r, key, db, cls)).rstrip()}")
            lines.append('| ' + ' | '.join([cls, kind] + cells) + ' |')
        targets = [(db, TARGET_TABLE.get((db, cls))) for db in dbs]
        if any(t for _, t in targets):
            cells = [f"D {t['D'][0]}/{t['D'][1]} · E {t['E'][0]}/{t['E'][1]}" if t else ''
                     for _, t in targets]
            lines.append('| ' + ' | '.join([cls, 'target'] + cells) + ' |')
    recs = ', '.join(f"{db} {next(by[(db, c)]['records'] for c in classes if (db, c) in by)}"
                     for db in dbs)
    lines.append(f"\nrecords: {recs}. `*` reaches the product target (Se and +P), `!` does not, "
                 "`-` no reference episodes of that class.")
    return '\n'.join(lines)


def beat_table(rows, tag=None):
    dbs = order([r['db'] for r in rows], DB_ORDER)
    by = {r['db']: r for r in rows}
    head = ['db', 'records', 'Q Se', 'Q +P', 'V Se', 'V +P', 'V F1', 'S Se', 'S +P', 'S F1',
            'QRS', 'VEB', 'SVEB']
    lines = []
    if tag:
        lines.append(f"**{tag}**\n")
    lines += ['| ' + ' | '.join(head) + ' |', '|' + '---|' * len(head)]
    for db in dbs:
        r = by[db]
        lines.append('| ' + ' | '.join(str(r.get(k, '-')) for k in [
            'db', 'records', 'Q_Se', 'Q_+P', 'V_Se', 'V_+P', 'V_F1', 'S_Se', 'S_+P', 'S_F1',
            'total_QRS', 'total_VEB', 'total_SVEB']) + ' |')
    return '\n'.join(lines)


def compare_table(results, kind):
    """Several result folders side by side: one column per folder, one row per db x class."""
    tags = list(results)
    if kind == 'rhythm':
        keys = order([(r['db'], r['class']) for t in tags for r in results[t]['rhythm']],
                     [(d, c) for d in DB_ORDER for c in CLASS_ORDER])
        metrics = ['D_F1', 'E_F1']
        get = lambda t, k: {(r['db'], r['class']): r for r in results[t]['rhythm']}.get(k)  # noqa
        label = lambda k: f"{k[0]} {k[1]}"  # noqa: E731
    else:
        keys = order([r['db'] for t in tags for r in results[t]['beats']], DB_ORDER)
        metrics = ['Q_Se', 'Q_+P', 'V_F1', 'S_F1']
        get = lambda t, k: {r['db']: r for r in results[t]['beats']}.get(k)  # noqa: E731
        label = lambda k: k  # noqa: E731
    head = ['cell'] + [f"{t} {m}" for t in tags for m in metrics]
    lines = ['| ' + ' | '.join(head) + ' |', '|' + '---|' * len(head)]
    for k in keys:
        cells = []
        for t in tags:
            r = get(t, k)
            cells += [str(r.get(m, '-')) if r else '' for m in metrics]
        lines.append('| ' + ' | '.join([label(k)] + cells) + ' |')
    return '\n'.join(lines)


def worst_records(per_record, n=8):
    """Records with the most false-positive / missed rhythm seconds, per class."""
    out = []
    for cls in order([r['class'] for r in per_record], CLASS_ORDER):
        rs = [r for r in per_record if r['class'] == cls]
        fp = sorted(rs, key=lambda r: -(num(r['test_duration_s']) or 0) * (1 - (num(r['D_+P']) or 0) / 100))
        fn = sorted(rs, key=lambda r: -(num(r['ref_duration_s']) or 0) * (1 - (num(r['D_Se']) or 100) / 100))
        out.append(f"{cls}: most FP seconds -> " + ', '.join(
            f"{r['db']}/{r['record']} (+{(num(r['test_duration_s']) or 0) * (1 - (num(r['D_+P']) or 0) / 100):.0f}s)"
            for r in fp[:n] if (num(r['D_+P']) or 0) < 100 and (num(r['test_duration_s']) or 0) > 0))
        out.append(f"{cls}: most missed seconds -> " + ', '.join(
            f"{r['db']}/{r['record']} (-{(num(r['ref_duration_s']) or 0) * (1 - (num(r['D_Se']) or 100) / 100):.0f}s)"
            for r in fn[:n] if (num(r['D_Se']) or 100) < 100))
    return '\n'.join(out)


# ----------------------------------------------------------------------------- output
def write_csv(rows, columns, path):
    with open(path, 'w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=columns, extrasaction='ignore')
        w.writeheader()
        w.writerows(rows)
    return path


def write_xlsx(sheets, path):
    """sheets = [(name, columns, rows)]; numbers stay numbers, '-' stays text."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        print(f"openpyxl not installed - {os.path.basename(path)} skipped")
        return None
    wb = Workbook()
    wb.remove(wb.active)
    for name, columns, rows in sheets:
        ws = wb.create_sheet(name[:31])
        ws.append(columns)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([num(r.get(c)) if num(r.get(c)) is not None else r.get(c, '-')
                       for c in columns])
        ws.freeze_panes = 'A2'
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(
                8, min(16, max(len(str(c.value or '')) for c in col) + 2))
    wb.save(path)
    return path


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    ap.add_argument('dirs', nargs='*', default=[os.path.join(HERE, 'ec57_results')],
                    help='EC57 result folder(s) (default: ec57_results/ next to this script)')
    ap.add_argument('--out', help='output folder (default: the first result folder)')
    ap.add_argument('--records', action='store_true',
                    help='also write per-record CSVs / sheets and print the worst records')
    ap.add_argument('--quiet', action='store_true', help='write files only')
    args = ap.parse_args(argv)

    results = {}
    for d in args.dirs:
        d = os.path.abspath(d.rstrip('/'))
        if not os.path.isdir(d):
            print(f"not a folder: {d}", file=sys.stderr)
            continue
        rh, rh_rec = rhythm_rows(d)
        bt, bt_rec = beat_rows(d)
        if not rh and not bt:
            print(f"no *_report_line.out under {d}", file=sys.stderr)
            continue
        results[os.path.basename(d)] = {'dir': d, 'rhythm': rh, 'rhythm_records': rh_rec,
                                        'beats': bt, 'beats_records': bt_rec}
    if not results:
        return 1

    out = os.path.abspath(args.out) if args.out else next(iter(results.values()))['dir']
    os.makedirs(out, exist_ok=True)
    md = []
    for tag, res in results.items():
        md.append(f"# EC57 summary - {tag}\n\nsource: `{res['dir']}`\n")
        if res['rhythm']:
            md.append("## Rhythm (epicmp)\n\n" + rhythm_table(res['rhythm']) + '\n')
        if res['beats']:
            md.append("## Beats (bxb)\n\n" + beat_table(res['beats']) + '\n')
        if args.records and res['rhythm_records']:
            md.append("## Worst records (duration)\n\n```\n" + worst_records(res['rhythm_records'])
                      + "\n```\n")
    if len(results) > 1:
        md.append("# Comparison\n\n## Rhythm F1 (duration / episode)\n\n"
                  + compare_table(results, 'rhythm') + "\n\n## Beats\n\n"
                  + compare_table(results, 'beats') + '\n')

    multi = len(results) > 1
    written = []
    for tag, res in results.items():
        suffix = f"_{tag}" if multi else ''
        if res['rhythm']:
            written.append(write_csv(res['rhythm'], RHYTHM_COLUMNS,
                                     os.path.join(out, f"ec57_summary{suffix}.csv")))
        if res['beats']:
            written.append(write_csv(res['beats'], BEAT_COLUMNS,
                                     os.path.join(out, f"ec57_beats_summary{suffix}.csv")))
        sheets = [('rhythm', RHYTHM_COLUMNS, res['rhythm']), ('beats', BEAT_COLUMNS, res['beats'])]
        if args.records:
            rec_cols = ['db', 'class', 'record', 'TPs', 'FN', 'TPp', 'FP', 'E_Se', 'E_+P', 'E_F1',
                        'D_Se', 'D_+P', 'D_F1', 'ref_duration_s', 'test_duration_s']
            brec_cols = ['db', 'record'] + BXB_COUNTS + BEAT_METRICS + ['V_F1', 'S_F1', 'RR_err']
            if res['rhythm_records']:
                written.append(write_csv(res['rhythm_records'], rec_cols,
                                         os.path.join(out, f"ec57_records_rhythm{suffix}.csv")))
                sheets.append(('rhythm_records', rec_cols, res['rhythm_records']))
            if res['beats_records']:
                written.append(write_csv(res['beats_records'], brec_cols,
                                         os.path.join(out, f"ec57_records_beats{suffix}.csv")))
                sheets.append(('beats_records', brec_cols, res['beats_records']))
        p = write_xlsx([s for s in sheets if s[2]], os.path.join(out, f"ec57_summary{suffix}.xlsx"))
        if p:
            written.append(p)
    md_path = os.path.join(out, 'ec57_summary.md')
    with open(md_path, 'w') as f:
        f.write('\n'.join(md))
    written.append(md_path)

    if not args.quiet:
        print('\n'.join(md))
        print("written:\n  " + '\n  '.join(written))
    return 0


if __name__ == '__main__':
    sys.exit(main())

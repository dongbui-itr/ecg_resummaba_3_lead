"""build_summary_xlsx.py <out.xlsx> <key> [key ...] : EC57 (rhythm + beats, all DBs) + dataset-eval v4 (beat mark/whole,
rhythm epicmp/event/sample) for each key from summary_111026/strip_eval/jobs/ec57_<key>.log and ec57/<key>/dataset_eval_v4/result.json,
plus a side-by-side comparison sheet."""
import json, re, sys, os, pandas as pd
J='/media/Project/ECG/Model_Dong/ecgr_rhythm/summary_111026/strip_eval/jobs'; E='/media/Project/ECG/Model_Dong/ecgr_rhythm/summary_111026/ec57'
out, keys = sys.argv[1], sys.argv[2:]
def parse(k):
    rh=[];bt=[];f1=[]
    for l in open(f'{J}/ec57_{k}.log'):
        m=re.match(r'\s+(\w+)\s+(AFIB|SVT|VT|AVB)\s+records=\s*(\d+)\s+E_Se=(\S+)\s+E_\+P=(\S+)\s+E_F1=(\S+)\s+D_Se=(\S+)\s+D_\+P=(\S+)\s+D_F1=(\S+)',l)
        if m: rh.append(dict(zip(['db','class','records','E_Se','E_+P','E_F1','D_Se','D_+P','D_F1'],m.groups())))
        m=re.match(r'\s+beats (\w+): Q_Se=(\S+) Q_\+P=(\S+) V_Se=(\S+) V_\+P=(\S+) S_Se=(\S+) S_\+P=(\S+) total_QRS=(\d+) total_VEB=(\d+) total_SVEB=(\d+) records=(\d+)',l)
        if m: bt.append(dict(zip(['db','Q_Se','Q_+P','V_Se','V_+P','S_Se','S_+P','total_QRS','total_VEB','total_SVEB','records'],m.groups())))
        m=re.search(r'macro F1.*: ([0-9.]+)',l)
        if m: f1.append(float(m.group(1)))
    return rh,bt,f1
D={k:parse(k) for k in keys}; V={}
for k in keys:
    p=f'{E}/{k}/dataset_eval_v4/result.json'; V[k]=json.load(open(p)) if os.path.exists(p) else None
with pd.ExcelWriter(out, engine='openpyxl') as xw:
    W=lambda df,name,index=False: df.to_excel(xw,sheet_name=name[:31],index=index)
    W(pd.DataFrame([[k, f'{J}/ec57_{k}.log', 'overlap 1 s, TTA id/flip/swap, AF gate', f"test macro F1 {min(D[k][2]):.2f}-{max(D[k][2]):.2f}" if D[k][2] else ''] for k in keys],columns=['key','ec57 log','inference','test split']),'info')
    rows=[]
    for k in keys:
        for r in D[k][0]: rows.append(dict(key=k,**r))
    W(pd.DataFrame(rows),'EC57_rhythm')
    rows=[]
    for k in keys:
        for r in D[k][1]: rows.append(dict(key=k,**r))
    W(pd.DataFrame(rows),'EC57_beats')
    # side by side: beats
    rows=[]
    for db in ['mitdb','nstdb','escdb','ahadb','afdb']:
        for m in ['Q_Se','Q_+P','V_Se','V_+P','S_Se','S_+P']:
            row=dict(db=db,metric=m)
            for k in keys:
                b={r['db']:r for r in D[k][1]}.get(db); row[k]=b[m] if b else ''
            rows.append(row)
    W(pd.DataFrame(rows),'beats_side_by_side')
    rows=[]
    for (db,c) in [('mitdb','AFIB'),('mitdb','SVT'),('mitdb','VT'),('mitdb','AVB'),('afdb','AFIB'),('escdb','VT'),('escdb','SVT')]:
        for m in ['E_Se','E_+P','E_F1','D_Se','D_+P','D_F1']:
            row=dict(db=db,cls=c,metric=m)
            for k in keys:
                r={(x['db'],x['class']):x for x in D[k][0]}.get((db,c)); row[k]=r[m] if r else ''
            rows.append(row)
    W(pd.DataFrame(rows),'rhythm_side_by_side')
    rows=[]
    for k in keys:
        if not V[k]: continue
        b=V[k]['beats']['bxb']
        rows.append(dict(key=k,scope='mark window',**b['mark_window'])); rows.append(dict(key=k,scope='whole strip',**b.get('whole_strip',{})))
    W(pd.DataFrame(rows),'v4_beat')
    rows=[]
    for k in keys:
        if not V[k]: continue
        for c,v in V[k]['rhythm']['epicmp'].items(): rows.append(dict(key=k,cls=c,E_Se=v[0],**{'E_+P':v[1]},D_Se=v[2],**{'D_+P':v[3]}))
    W(pd.DataFrame(rows),'v4_rhythm_epicmp')
    for k in keys:
        if not V[k]: continue
        r=V[k]['rhythm']
        if 'event_cm' in r:
            cm=pd.DataFrame(r['event_cm']).T; cm['(none)']=pd.Series(r.get('event_none',{})); cm.index.name='true \\ pred'; W(cm,f'{k}_event_cm',True)
        if 'sample_cm' in r:
            names=['SINUS','AFIB','SVT','VT','AVB','NOISE'][:len(r['sample_cm'][0])]; sc=pd.DataFrame(r['sample_cm'],columns=names); sc.index=['SINUS','AFIB','SVT','VT','AVB2','AVB3'][:len(sc)]; sc.index.name='true \\ pred'; W(sc,f'{k}_sample_cm',True)
    for k in keys:
        for name,path in [(f'{k}_mitdb_AFIB',f'{E}/{k}/mitdb/mitdb_AFIB_report_line.out'),(f'{k}_mitdb_beats',f'{E}/{k}/beats/mitdb/mitdb_QRS_report_line.out')]:
            if not os.path.exists(path): continue
            L=[l for l in open(path).read().splitlines() if re.match(r'\s*(Record|\d+|Sum|Gross|Average)',l)]
            h=L[0].split(); data=[l.split() for l in L[1:]]; w=max(len(h),max(len(d) for d in data))
            W(pd.DataFrame([d+['']*(w-len(d)) for d in data],columns=h+[f'c{i}' for i in range(w-len(h))]),name)
from openpyxl import load_workbook
wb=load_workbook(out)
for ws in wb.worksheets:
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=max(10,min(60,max(len(str(c.value)) if c.value is not None else 0 for c in col)+2))
wb.save(out); print(out, wb.sheetnames)

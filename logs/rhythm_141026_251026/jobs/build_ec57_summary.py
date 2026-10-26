"""EC57 + rhythm_eval + data-eval summary workbook: base (official decode) vs tuned models.

Sheets: Rhythm (EC57 epicmp per database -> class -> model; rhythm_eval inside the reviewer mark
vs the 5.4.1 / event reference; test-split data eval), Beat (bxb per database -> model).
"""
import csv
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, '/media/Project/ECG/Source_Dong/project_ai/ecg_resummaba_3_lead')
from ecgr.rhythm import strip_eval as SE  # noqa: E402

W = '/media/Project/ECG/Model_Dong/ecgr_rhythm'
S = f'{W}/summary_111026'
T = f'{S}/strip_eval'
MODELS = [  # key, label, params, ec57 rhythm dir, beat dirs (in order, later wins), test eval dir, strip json
    ('base1m', 'Dual 1m e26 — gốc, decode cũ', '989k', f'{S}/ec57/dual1m_s1_e26',
     [f'{W}/release_dualunet_1m_e26/work/eval/ec57/ec57_261008_1122/beats', f'{S}/ec57/dual1m_s1_e26/beats'],
     f'{W}/091026_rhythm_dual_small/eval/teacher_1m_e26', (f'{T}/base_dual1m_s1_e26_rhythm_eval.json', None)),
    ('tuned1m', 'Dual 1m — fine-tune + decode tune', '989k', f'{S}/ec57/tuned_dual1m_ft2',
     [f'{S}/ec57/tuned_dual1m_ft2/beats'], f'{S}/eval/tuned_dual1m_ft2',
     (f'{T}/final_dual1m_ft2_best_v2.json', 'rhythm_eval')),
    ('H1m', 'Dual 1m — cách 2 (portal + Holter)', '989k', f'{S}/ec57/H_dual1m_ft2',
     [f'{S}/ec57/H_dual1m_ft2/beats'], f'{S}/eval/tuned_dual1m_ft2',
     (f'{T}/final_dual1m_ft2_best_H.json', 'rhythm_eval')),
    ('base100k', 'Dual 100k — gốc, decode cũ', '80k', f'{S}/ec57/dual100k',
     [f'{W}/091026_rhythm_dual_small/ec57/ec57_dual_100k_official/beats', f'{S}/ec57/dual100k/beats'],
     f'{W}/091026_rhythm_dual_small/eval/dual_100k', (f'{T}/base_dual100k_rhythm_eval.json', None)),
    ('tuned100k', 'Dual 100k — fine-tune + decode tune', '80k', f'{S}/ec57/tuned_dual100k_ft2',
     [f'{S}/ec57/tuned_dual100k_ft2/beats'], f'{S}/eval/tuned_dual100k_ft2',
     (f'{T}/final_dual100k_ft2_last_v2.json', 'rhythm_eval')),
    ('H100k', 'Dual 100k — cách 2 (portal + Holter)', '80k', f'{S}/ec57/H_dual100k',
     [f'{S}/ec57/H_dual100k/beats'], f'{W}/091026_rhythm_dual_small/eval/dual_100k',
     (f'{T}/final_dual100k_H.json', 'rhythm_eval')),
    ('base30k', 'Dual 30k — gốc, decode cũ', '23k', f'{S}/ec57/dual30k',
     [f'{W}/091026_rhythm_dual_small/ec57/ec57_dual_30k_official/beats', f'{S}/ec57/dual30k/beats'],
     f'{W}/091026_rhythm_dual_small/eval/dual_30k', (f'{T}/base_dual30k_rhythm_eval.json', None)),
    ('tuned30k', 'Dual 30k — fine-tune + decode tune', '23k', f'{S}/ec57/tuned_dual30k_ft',
     [f'{S}/ec57/tuned_dual30k_ft/beats'], f'{S}/eval/tuned_dual30k_ft',
     (f'{T}/final_dual30k_ft_e3_v2.json', 'rhythm_eval')),
    ('H30k', 'Dual 30k — cách 2 (portal + Holter)', '23k', f'{S}/ec57/H_dual30k',
     [f'{S}/ec57/H_dual30k/beats'], f'{W}/091026_rhythm_dual_small/eval/dual_30k',
     (f'{T}/final_dual30k_H.json', 'rhythm_eval')),
]
EC57_CELLS = [('mitdb', ['AFIB', 'SVT', 'VT', 'AVB']), ('afdb', ['AFIB']), ('escdb', ['SVT', 'VT'])]
EC57_TARGET = {('mitdb', 'AFIB'): (98, 90, 80, 92), ('afdb', 'AFIB'): (96, 99, 81, 93),
               ('mitdb', 'SVT'): (61, 24, 62, 19), ('mitdb', 'VT'): (75, 57, 82, 58),
               ('mitdb', 'AVB'): (98, 87, 100, 100)}   # D Se, D +P, E Se, E +P ("rhythm 3.0.6")
BEAT_DBS = ['mitdb', 'escdb', 'afdb', 'nstdb', 'ahadb']
SNRS = [('clean', 'Sạch'), ('snr_18', '18 dB'), ('snr_12', '12 dB'), ('snr_6', '6 dB'),
        ('snr_0', '0 dB'), ('snr_-6', '−6 dB')]


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def read_csv(path, key):
    return {key(r): r for r in csv.DictReader(open(path))} if os.path.exists(path) else {}


def strip_metrics(spec):
    path, sub = spec
    d = json.load(open(path))
    return d[sub] if sub else d


def test_eval(d):
    j = [f for f in os.listdir(d) if f.endswith('_test_eval.json')] if os.path.isdir(d) else []
    return json.load(open(os.path.join(d, j[0]))) if j else None


F = lambda **k: Font(**{'name': 'Arial', 'size': 10, **k})  # noqa: E731
HDR = PatternFill('solid', fgColor='1F3864')
GRP = PatternFill('solid', fgColor='F2F2F2')
TUN = PatternFill('solid', fgColor='EAF1FB')
HOL = PatternFill('solid', fgColor='FFF2CC')
BEST = PatternFill('solid', fgColor='C6EFCE')
FAIL = PatternFill('solid', fgColor='FFC7CE')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
TOP = Border(left=thin, right=thin, bottom=thin, top=Side(style='medium', color='1F3864'))
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
GREEN = Font(name='Arial', size=10, bold=True, color='006100')


def header(ws, r, labels):
    for i, t in enumerate(labels, 1):
        c = ws.cell(r, i, t)
        c.font = F(bold=True, color='FFFFFF')
        c.fill = HDR
        c.alignment = CENTER
        c.border = BOX


def put(ws, r, c, v, fmt=None, fill=None, **font):
    cell = ws.cell(r, c, v)
    cell.font = F(**font)
    cell.border = BOX
    if fill:
        cell.fill = fill
    if isinstance(v, (int, float)) or (isinstance(v, str) and v.startswith('=')):
        cell.alignment = Alignment(horizontal='center')
        if fmt:
            cell.number_format = fmt
    return cell


def f1(r, a, b):
    x, y = f'{L(a)}{r}', f'{L(b)}{r}'
    return f'=IF(AND(ISNUMBER({x}),ISNUMBER({y})),IF({x}+{y}>0,2*{x}*{y}/({x}+{y}),0),"-")'


def ok(r, col, target_col):
    return f'=IF(ISNUMBER({col}{r}),IF({col}{r}>={target_col}{r},"✓","✗"),"-")'


def title(ws, r, text, note=None, width='P'):
    ws.cell(r, 1, text).font = F(bold=True, size=12)
    if note:
        ws.cell(r + 1, 1, note).font = F(italic=True, color='595959')
        ws.merge_cells(f'A{r + 1}:{width}{r + 1}')
        ws.cell(r + 1, 1).alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r + 1].height = 44
        return r + 3
    return r + 2


def best_rule(ws, col, r0, r1):
    ws.conditional_formatting.add(f'{col}{r0}:{col}{r1}', FormulaRule(
        formula=[f'AND(ISNUMBER({col}{r0}),{col}{r0}=MAX(${col}${r0}:${col}${r1}))'],
        fill=BEST, font=GREEN))


def check_rules(ws, rng, first):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="✗"'], fill=FAIL))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="✓"'], fill=BEST))


def main(out):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Rhythm'
    widths = [12, 7, 36, 9, 10, 10, 10, 10, 10, 10, 10, 10, 8, 8, 9, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[L(i)].width = w

    # ---- A. EC57 -------------------------------------------------------------------------
    r = title(ws, 1, 'A. EC57 — epicmp (F1 duration / episode), theo database → lớp → model',
              'Model "gốc" = checkpoint distill/release với decode chính thức (config mặc định). Model "tune" = '
              'fine-tune trên 5 nguồn portal của rhythm_eval + decode/post-process chọn trên validation strip_val2 '
              '(split eval, không trùng study với rhythm_eval). Hop 5 s, TTA id,flip,swap, AFL = AF (epicmp -x). '
              'Model "cách 2" = một decode chọn trên strip_val2 VÀ không tụt dưới mức của model gốc trên Holter validation (eval side '
              'ltafdb/nsrdb/incartdb/svdb); chọn giữa bản gốc và bản fine-tune theo validation (1m: fine-tune; 100k, 30k: gốc). '
              'Target = bảng sản phẩm "rhythm 3.0.6". nstdb / ahadb không có cơn rhythm tham chiếu (chỉ có ở sheet Beat). '
              'Nền xanh = tune portal, nền vàng = cách 2.')
    cols = ['Database', 'Lớp', 'Model', 'Tham số', 'D Se', 'D +P', 'D F1', 'E Se', 'E +P', 'E F1',
            'Target D F1', 'Target E F1', 'Đạt D', 'Đạt E', 'Records']
    header(ws, r, cols)
    r += 1
    rhy = {m[0]: read_csv(f'{m[3]}/rhythm_ec57_summary.csv', lambda x: (x['db'], x['class'])) for m in MODELS}
    for db, classes in EC57_CELLS:
        for cls in classes:
            first = r
            for key, label, params, *_ in MODELS:
                row = rhy[key].get((db, cls), {})
                fill = TUN if key.startswith('tuned') else (HOL if key.startswith('H') else None)
                put(ws, r, 1, db, bold=r == first, fill=GRP)
                put(ws, r, 2, cls, bold=r == first, fill=GRP)
                put(ws, r, 3, label, fill=fill)
                put(ws, r, 4, params, fill=fill)
                for c, k in ((5, 'D_Se'), (6, 'D_+P'), (8, 'E_Se'), (9, 'E_+P')):
                    v = num(row.get(k))
                    put(ws, r, c, v if v is not None else '-', '0')
                put(ws, r, 7, f1(r, 5, 6), '0.0', bold=True)
                put(ws, r, 10, f1(r, 8, 9), '0.0', bold=True)
                t = EC57_TARGET.get((db, cls))
                if t:
                    put(ws, r, 11, round(2 * t[0] * t[1] / (t[0] + t[1]), 1), '0.0', color='0000FF')
                    put(ws, r, 12, round(2 * t[2] * t[3] / (t[2] + t[3]), 1), '0.0', color='0000FF')
                    put(ws, r, 13, ok(r, 'G', 'K'))
                    put(ws, r, 14, ok(r, 'J', 'L'))
                    if r == first:
                        ws.cell(r, 11).comment = Comment(
                            f'Target rhythm 3.0.6: D Se/+P {t[0]}/{t[1]}, E Se/+P {t[2]}/{t[3]}', 'summary')
                else:
                    for c in (11, 12, 13, 14):
                        put(ws, r, c, '-')
                put(ws, r, 15, num(row.get('records')), '0')
                r += 1
            for col in ('G', 'J'):
                best_rule(ws, col, first, r - 1)
            for c in range(1, len(cols) + 1):
                ws.cell(first, c).border = TOP
    check_rules(ws, f'M5:N{r - 1}', 'M5')

    # ---- B. rhythm_eval ------------------------------------------------------------------
    r += 1
    r = title(ws, r, 'B. rhythm_eval (holdout portal, 4.909 sự kiện) — chấm bên trong cửa sổ bác sĩ duyệt, so với bảng tham khảo',
              'Giao thức ecgr/rhythm/strip_eval.py: chỉ cửa sổ đánh dấu (10-50 s) có nhãn nên chỉ chấm trong đó. '
              'B1 = epicmp segment-to-segment (target "Version 5.4.1"); B2 = sample-to-sample recall (đường chéo 5.4.1); '
              'B3 = event-level Se/PPV/F1 (hình tham khảo thứ 2). AVB = AVB2+AVB3 gộp, phải vượt hàng tốt hơn của AVB2/AVB3.')
    strip = {m[0]: strip_metrics(m[6]) for m in MODELS}
    header(ws, r, ['Database', 'Lớp', 'Model', 'Tham số', 'E Se', 'E +P', 'D Se', 'D +P',
                   'Tgt E Se', 'Tgt E +P', 'Tgt D Se', 'Tgt D +P', 'Đạt (4 ô)', '', '', ''])
    r += 1
    b0 = r
    for cls in SE.EVENT_CLASSES:
        first = r
        t = SE.TARGET_EPICMP[cls]
        for key, label, params, *_ in MODELS:
            v = strip[key]['epicmp'][cls]
            fill = TUN if key.startswith('tuned') else (HOL if key.startswith('H') else None)
            put(ws, r, 1, 'rhythm_eval', bold=r == first, fill=GRP)
            put(ws, r, 2, cls, bold=r == first, fill=GRP)
            put(ws, r, 3, label, fill=fill)
            put(ws, r, 4, params, fill=fill)
            for i in range(4):
                put(ws, r, 5 + i, round(v[i], 1), '0.0', bold=True)
                put(ws, r, 9 + i, t[i], '0', color='0000FF')
            put(ws, r, 13, f'=IF(AND(E{r}>=I{r},F{r}>=J{r},G{r}>=K{r},H{r}>=L{r}),"✓","✗")')
            r += 1
        for col in ('E', 'F', 'G', 'H'):
            best_rule(ws, col, first, r - 1)
        for c in range(1, 14):
            ws.cell(first, c).border = TOP
    for col, tcol in (('E', 'I'), ('F', 'J'), ('G', 'K'), ('H', 'L')):
        ws.conditional_formatting.add(f'{col}{b0}:{col}{r - 1}', FormulaRule(formula=[f'{col}{b0}<{tcol}{b0}'], fill=FAIL))
    check_rules(ws, f'M{b0}:M{r - 1}', f'M{b0}')

    r += 1
    header(ws, r, ['Database', 'Hàng', 'Model', 'Tham số', 'Recall %', 'Target', 'Đạt', '', '', '', '', '', '', '', '', ''])
    r += 1
    b2 = r
    for t in SE.REF_TYPES:
        first = r
        for key, label, params, *_ in MODELS:
            fill = TUN if key.startswith('tuned') else (HOL if key.startswith('H') else None)
            put(ws, r, 1, 'rhythm_eval', bold=r == first, fill=GRP)
            put(ws, r, 2, t, bold=r == first, fill=GRP)
            put(ws, r, 3, label, fill=fill)
            put(ws, r, 4, params, fill=fill)
            put(ws, r, 5, round(strip[key]['sample_recall'][t], 2), '0.00', bold=True)
            put(ws, r, 6, SE.TARGET_SAMPLE_RECALL[t], '0.00', color='0000FF')
            put(ws, r, 7, ok(r, 'E', 'F'))
            r += 1
        best_rule(ws, 'E', first, r - 1)
        for c in range(1, 8):
            ws.cell(first, c).border = TOP
    check_rules(ws, f'G{b2}:G{r - 1}', f'G{b2}')

    r += 1
    header(ws, r, ['Database', 'Lớp', 'Model', 'Tham số', 'Se', 'PPV', 'F1', 'Tgt Se', 'Tgt PPV', 'Tgt F1',
                   'Đạt F1', '', '', '', '', ''])
    r += 1
    b3 = r
    for cls in SE.EVENT_CLASSES:
        first = r
        t = SE.TARGET_EVENT[cls]
        for key, label, params, *_ in MODELS:
            se, ppv, _f = strip[key]['event'][cls]
            fill = TUN if key.startswith('tuned') else (HOL if key.startswith('H') else None)
            put(ws, r, 1, 'rhythm_eval', bold=r == first, fill=GRP)
            put(ws, r, 2, cls, bold=r == first, fill=GRP)
            put(ws, r, 3, label, fill=fill)
            put(ws, r, 4, params, fill=fill)
            put(ws, r, 5, round(se, 1), '0.0')
            put(ws, r, 6, round(ppv, 1), '0.0')
            put(ws, r, 7, f1(r, 5, 6), '0.0', bold=True)
            for i in range(3):
                put(ws, r, 8 + i, t[i], '0.0', color='0000FF')
            put(ws, r, 11, ok(r, 'G', 'J'))
            r += 1
        best_rule(ws, 'G', first, r - 1)
        for c in range(1, 12):
            ws.cell(first, c).border = TOP
    check_rules(ws, f'K{b3}:K{r - 1}', f'K{b3}')
    r += 1
    ws.cell(r, 1, 'Tổng ô target rhythm_eval đạt (36 ô: B1 16 + B2 6 + B3 12 + 2 hàng AVB2/AVB3 Se):').font = F(bold=True)
    r += 1
    for key, label, *_ in MODELS:
        met = sum(x[5] for x in SE.compare(strip[key]))
        put(ws, r, 3, label, fill=TUN if key.startswith('tuned') else (HOL if key.startswith('H') else None))
        put(ws, r, 4, f'{met}/36', bold=True)
        r += 1

    # ---- C. data eval -------------------------------------------------------------------
    r += 1
    r = title(ws, r, 'C. Data eval — tập test của dữ liệu train (7.233 window 10 s), F1 macro theo giây',
              'Nhiễu thêm vào mọi window với SNR cố định; F1 theo lớp ở tập sạch. Không có decode/post-process (đầu ra thô theo giây).')
    header(ws, r, ['Database', 'Model', 'Tham số'] + [s for _, s in SNRS] + ['SINUS', 'AFIB', 'SVT', 'VT', 'AVB', '', ''])
    r += 1
    c0 = r
    for key, label, params, _a, _b, tdir, _s in MODELS:
        t = test_eval(tdir)
        fill = TUN if key.startswith('tuned') else (HOL if key.startswith('H') else None)
        put(ws, r, 1, 'test split', fill=GRP)
        put(ws, r, 2, label, fill=fill)
        put(ws, r, 3, params, fill=fill)
        for i, (k, _) in enumerate(SNRS):
            put(ws, r, 4 + i, round(100 * t[k]['macro_f1'], 2) if t else '-', '0.00')
        pc = {x['cls']: x for x in t['clean']['per_class']} if t else {}
        for i, cls in enumerate(['SINUS', 'AFIB', 'SVT', 'VT', 'AVB']):
            put(ws, r, 10 + i, round(100 * pc[cls]['f1'], 1) if cls in pc else '-', '0.0')
        r += 1
    for i in range(11):
        best_rule(ws, L(4 + i), c0, r - 1)
    ws.column_dimensions['B'].width = 9
    ws.freeze_panes = 'E5'

    # ---- Beat ---------------------------------------------------------------------------
    wb2 = wb.create_sheet('Beat')
    ws = wb2
    for i, w in enumerate([10, 36, 9, 9, 9, 9, 9, 9, 9, 9, 9, 9, 11, 10, 10, 9], 1):
        ws.column_dimensions[L(i)].width = w
    r = title(ws, 1, 'BEAT — bxb (AAMI EC57), theo database → model',
              'Beat hypothesis = ngõ ra beat của model sau decode (model "tune": beat stage tắt nên ký hiệu beat là đầu ra thô '
              'của decoder beat; model "gốc": sau beat post-processing chính thức). afdb chấm trên .qrs (không có loại beat). '
              'F1 = công thức từ Se và +P. Mục tiêu beat (README): mitdb QRS ≥ 99.95, SVEB Se > 43, SVEB +P > 80.')
    cb = ['Database', 'Model', 'Tham số', 'QRS Se', 'QRS +P', 'QRS F1', 'VEB Se', 'VEB +P', 'VEB F1',
          'SVEB Se', 'SVEB +P', 'SVEB F1', 'Tổng QRS', 'Tổng VEB', 'Tổng SVEB', 'Records']
    header(ws, r, cb)
    r += 1
    beats = {}
    for key, *_rest in MODELS:
        d = {}
        for p in _rest[3]:
            for row in csv.DictReader(open(f'{p}/ec57_summary.csv')) if os.path.exists(f'{p}/ec57_summary.csv') else []:
                d[row['db']] = row
        beats[key] = d
    for db in BEAT_DBS:
        first = r
        for key, label, params, *_ in MODELS:
            row = beats[key].get(db)
            if row is None:
                continue
            fill = TUN if key.startswith('tuned') else (HOL if key.startswith('H') else None)
            put(ws, r, 1, db, bold=r == first, fill=GRP)
            put(ws, r, 2, label, fill=fill)
            put(ws, r, 3, params, fill=fill)
            for c, k in ((4, 'Q_Se'), (5, 'Q_+P'), (7, 'V_Se'), (8, 'V_+P'), (10, 'S_Se'), (11, 'S_+P')):
                v = num(row.get(k))
                put(ws, r, c, v if v is not None else '-', '0.00')
            for c, a, b in ((6, 4, 5), (9, 7, 8), (12, 10, 11)):
                put(ws, r, c, f1(r, a, b), '0.00', bold=True)
            for c, k in ((13, 'total_QRS'), (14, 'total_VEB'), (15, 'total_SVEB'), (16, 'records')):
                put(ws, r, c, num(row.get(k)), '#,##0')
            r += 1
        if r == first:
            continue
        for col in ('F', 'I', 'L'):
            best_rule(ws, col, first, r - 1)
        for c in range(1, len(cb) + 1):
            ws.cell(first, c).border = TOP
    ws.freeze_panes = 'D4'
    wb.save(out)
    print(out)


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else f'{S}/ec57_summary_tuned_vs_base.xlsx')

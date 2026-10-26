"""build_dataset_v4_summary.py key ... -> summary_111026/dataset_eval_v4_summary.xlsx"""
import json, sys
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter as L
S = '/media/Project/ECG/Model_Dong/ecgr_rhythm/summary_111026'
LABEL = {'v4_scratch_best': 'v4 (train từ đầu, data v4)', 'dual1m_s1_e26': 'Gốc 1m e26'}
keys = sys.argv[1:]; R = {k: json.load(open(f'{S}/ec57/{k}/dataset_eval_v4/result.json')) for k in keys}
HD = Font(name='Arial', bold=True, color='FFFFFF'); HF = PatternFill('solid', start_color='1F4E78')
DG = PatternFill('solid', start_color='E2EFDA'); TF = PatternFill('solid', start_color='D9E1F2')
th = Side(style='thin', color='BFBFBF'); BX = Border(left=th, right=th, top=th, bottom=th)
def title(ws, r, t, sub=None):
    ws.cell(r, 1, t).font = Font(name='Arial', bold=True, size=12, color='1F4E78')
    if sub: ws.cell(r + 1, 1, sub).font = Font(name='Arial', italic=True, color='595959')
    return r + (3 if sub else 2)
def head(ws, r, c, hs):
    for j, h in enumerate(hs):
        x = ws.cell(r, c + j, h); x.font = HD; x.fill = HF; x.border = BX; x.alignment = Alignment(horizontal='center', wrap_text=True)
def put(ws, r, c, v, fmt='0.00', b=False, f=None):
    x = ws.cell(r, c, v); x.number_format = fmt; x.font = Font(name='Arial', bold=b); x.border = BX
    if f: x.fill = f
def num(v):
    try: return float(v)
    except (TypeError, ValueError): return None
wb = Workbook(); ws = wb.active; ws.title = 'Summary bxb + epicmp'
r = title(ws, 1, 'dataset-eval/v4 - bảng tổng hợp bxb (beat) và epicmp (rhythm)',
          'Pipeline triển khai: inference hop 5 s + TTA, decode chuẩn (official) gồm post-process beat. beat-eval-dataset 5,227 strip (tham chiếu .atr); rhythm-eval-dataset 4,678 event (chấm trong vùng đánh dấu).')
ws.cell(r, 1, 'A. Beat - bxb (Gross)').font = Font(name='Arial', bold=True); r += 1
BK = [('Q_Se', 'QRS Se'), ('Q_+P', 'QRS +P'), ('V_Se', 'VEB Se'), ('V_+P', 'VEB +P'), ('S_Se', 'SVEB Se'), ('S_+P', 'SVEB +P')]
head(ws, r, 1, ['Phạm vi', 'Model'] + [b for _, b in BK]); r += 1
for prot, pl in (('mark_window', 'Vùng đánh dấu (chuẩn beat team)'), ('whole_strip', 'Toàn bộ strip 60 s')):
    for k in keys:
        row = R[k]['beats']['bxb'].get(prot, {})
        put(ws, r, 1, pl, '@'); put(ws, r, 2, LABEL.get(k, k), '@', True)
        for j, (kk, _) in enumerate(BK): put(ws, r, 3 + j, num(row.get(kk)))
        r += 1
r += 1; ws.cell(r, 1, 'B. Rhythm - epicmp (Gross, trong vùng đánh dấu)').font = Font(name='Arial', bold=True); r += 1
head(ws, r, 1, ['Lớp', 'Model', 'E Se', 'E +P', 'E F1', 'D Se', 'D +P', 'D F1']); r += 1
TGT = {'AFIB': (90, 94, 88, 96), 'SVT': (89, 88, 60, 78), 'VT': (81, 90, 62, 86), 'AVB': (25, 77, 31, 65)}
f1 = lambda a, b: (2 * a * b / (a + b)) if a is not None and b is not None and (a + b) > 0 else None
for c in ['AFIB', 'SVT', 'VT', 'AVB']:
    for k in keys + ['5.4.1 (tham chiếu)']:
        v = TGT[c] if k.startswith('5.4.1') else tuple(R[k]['rhythm']['epicmp'][c])
        v = [None if (x is None or (isinstance(x, float) and np.isnan(x))) else x for x in v]
        put(ws, r, 1, c, '@', True); put(ws, r, 2, LABEL.get(k, k), '@', not k.startswith('5.4.1'), TF if k.startswith('5.4.1') else None)
        for j, x in enumerate([v[0], v[1], f1(v[0], v[1]), v[2], v[3], f1(v[2], v[3])]): put(ws, r, 3 + j, x)
        r += 1
r += 1; ws.cell(r, 1, 'C. Rhythm - event confusion (Se / PPV / F1 theo event)').font = Font(name='Arial', bold=True); r += 1
head(ws, r, 1, ['Lớp', 'Model', 'Se', 'PPV', 'F1']); r += 1
for c in ['AFIB', 'SVT', 'VT', 'AVB']:
    for k in keys:
        put(ws, r, 1, c, '@', True); put(ws, r, 2, LABEL.get(k, k), '@', True)
        for j, x in enumerate(R[k]['rhythm']['event_stats'][c]): put(ws, r, 3 + j, x)
        r += 1
ws.column_dimensions['A'].width = 30; ws.column_dimensions['B'].width = 26
for j in range(3, 10): ws.column_dimensions[L(j)].width = 10
# ---- confusion sheets: counts + %
def matrix_sheet(name, sub, rows_lbl, cols_lbl, get, extra=None):
    ws = wb.create_sheet(name); r = title(ws, 1, name, sub)
    for k in keys:
        cm = np.array(get(R[k]), float)
        for pct in (False, True):
            ws.cell(r, 1, f'{LABEL.get(k, k)} - {"tỷ lệ % theo hàng" if pct else "số lượng"}').font = Font(name='Arial', bold=True); r += 1
            head(ws, r, 1, ['Nhãn \\ Dự đoán'] + cols_lbl + ['Tổng hàng'] + (['Se %', 'P+ %'] if extra else [])); r += 1; top = r
            for i, rl in enumerate(rows_lbl):
                tot = cm[i].sum(); put(ws, r, 1, rl, '@', True)
                for j in range(len(cols_lbl)):
                    v = (cm[i, j] / tot * 100 if tot else None) if pct else int(cm[i, j])
                    put(ws, r, 2 + j, v, '0.00' if pct else '#,##0', extra and extra(i, j), DG if extra and extra(i, j) else None)
                put(ws, r, 2 + len(cols_lbl), int(tot), '#,##0', True)
                if extra:
                    d = [j for j in range(len(cols_lbl)) if extra(i, j)]
                    if d and tot:
                        jj = d[0]; colsum = cm[:, jj].sum()
                        put(ws, r, 3 + len(cols_lbl), cm[i, jj] / tot * 100, '0.00', True)
                        put(ws, r, 4 + len(cols_lbl), cm[i, jj] / colsum * 100 if colsum else None, '0.00', True)
                r += 1
            if pct: ws.conditional_formatting.add(f'B{top}:{L(1 + len(cols_lbl))}{r - 1}', ColorScaleRule(start_type='num', start_value=0, start_color='FFFFFF', end_type='num', end_value=100, end_color='63BE7B'))
            r += 1
    ws.column_dimensions['A'].width = 34
    for j in range(2, 4 + len(cols_lbl) + 2): ws.column_dimensions[L(j)].width = 11
matrix_sheet('Beat confusion', 'beat-eval-dataset, beats trong vùng đánh dấu, khớp 150 ms; hàng "-" = phát hiện thừa, cột "Bỏ sót" = beat tham chiếu không được phát hiện. Beat sau post-process (đúng file bxb chấm).',
             ['N', 'S', 'V', '- (phát hiện thừa)'], ['N', 'S', 'V', 'Bỏ sót'], lambda x: x['beats']['confusion'], lambda i, j: i == j and i < 3)
PN = ['SINUS', 'AFIB', 'SVT', 'VT', 'AVB', 'NOISE']; RT = ['SINUS', 'AFIB', 'SVT', 'VT', 'AVB2', 'AVB3']
matrix_sheet('Rhythm sample confusion', 'rhythm-eval-dataset, sample-to-sample trên lưới 40 ms trong vùng đánh dấu, sau decode chuẩn; AVB2/AVB3 đúng khi dự đoán AVB.',
             RT, PN, lambda x: x['rhythm']['sample_cm'], lambda i, j: PN[j] == ('AVB' if RT[i] in ('AVB2', 'AVB3') else RT[i]))
EC = ['AFIB', 'SVT', 'VT', 'AVB']
matrix_sheet('Rhythm event confusion', 'rhythm-eval-dataset, 1 event = 1 đơn vị: lớp dự đoán chiếm đa số trong vùng đánh dấu (lớp loạn nhịp; nếu không có thì lấy xác suất trung bình cao nhất).',
             ['AFIB', 'SVT', 'VT', 'AVB2', 'AVB3'], EC, lambda x: [[x['rhythm']['event_cm'][t][c] for c in EC] for t in ['AFIB', 'SVT', 'VT', 'AVB2', 'AVB3']],
             lambda i, j: EC[j] == ['AFIB', 'SVT', 'VT', 'AVB', 'AVB'][i])
out = f'{S}/dataset_eval_v4_summary.xlsx'; wb.save(out); print('saved', out)

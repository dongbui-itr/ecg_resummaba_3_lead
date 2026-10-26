"""build_ckpt_report.py <key> [base_key] -> summary_111026/ec57/<key>/<key>_report.xlsx
EC57 (mitdb nstdb ahadb escdb afdb: epicmp + bxb + confusion) and dataset-eval/v4 (beat, rhythm)."""
import json, sys, csv
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter as L
S = '/media/Project/ECG/Model_Dong/ecgr_rhythm/summary_111026/ec57'
key = sys.argv[1]; bk = sys.argv[2] if len(sys.argv) > 2 else None
DBS = ['mitdb', 'nstdb', 'ahadb', 'escdb', 'afdb']; C = ['SINUS', 'AFIB', 'SVT', 'VT', 'AVB']
HD = Font(name='Arial', bold=True, color='FFFFFF'); HF = PatternFill('solid', start_color='1F4E78')
DG = PatternFill('solid', start_color='E2EFDA'); TF = PatternFill('solid', start_color='D9E1F2'); GR = Font(name='Arial', color='808080')
th = Side(style='thin', color='BFBFBF'); BX = Border(left=th, right=th, top=th, bottom=th)
def rd(path):
    try: return list(csv.DictReader(open(path)))
    except OSError: return []
def num(v):
    try: return float(v)
    except (TypeError, ValueError): return None
def title(ws, r, t, sub=None):
    ws.cell(r, 1, t).font = Font(name='Arial', bold=True, size=12, color='1F4E78')
    if sub: ws.cell(r + 1, 1, sub).font = Font(name='Arial', italic=True, color='595959')
    return r + (3 if sub else 2)
def head(ws, r, hs, c=1):
    for j, h in enumerate(hs):
        x = ws.cell(r, c + j, h); x.font = HD; x.fill = HF; x.border = BX; x.alignment = Alignment(horizontal='center', wrap_text=True)
def put(ws, r, c, v, fmt='0.00', b=False, f=None, font=None):
    x = ws.cell(r, c, v); x.number_format = fmt; x.font = font or Font(name='Arial', bold=b); x.border = BX
    if f: x.fill = f
def widths(ws, a=26, rest=11, n=14):
    ws.column_dimensions['A'].width = a
    for j in range(2, n): ws.column_dimensions[L(j)].width = rest
def matrix(ws, r, label, cm, rows, cols, diag, total_label='Tổng hàng'):
    cm = np.array(cm, float)
    for pct in (False, True):
        ws.cell(r, 1, f'{label} - {"tỷ lệ % theo hàng" if pct else "số lượng"}').font = Font(name='Arial', bold=True); r += 1
        head(ws, r, ['Nhãn \\ Dự đoán'] + cols + [total_label, 'Se %', 'P+ %']); r += 1; top = r
        for i, rl in enumerate(rows):
            tot = cm[i].sum(); put(ws, r, 1, rl, '@', True)
            for j in range(len(cols)):
                v = (cm[i, j] / tot * 100 if tot else None) if pct else int(cm[i, j])
                d = diag(i, j); put(ws, r, 2 + j, v, '0.00' if pct else '#,##0', d, DG if d else None)
            put(ws, r, 2 + len(cols), int(tot), '#,##0', True)
            dj = [j for j in range(len(cols)) if diag(i, j)]
            if dj and tot:
                jj = dj[0]; cs = cm[:, jj].sum()
                put(ws, r, 3 + len(cols), cm[i, jj] / tot * 100, b=True); put(ws, r, 4 + len(cols), cm[i, jj] / cs * 100 if cs else None, b=True)
            r += 1
        if pct: ws.conditional_formatting.add(f'B{top}:{L(1 + len(cols))}{r - 1}', ColorScaleRule(start_type='num', start_value=0, start_color='FFFFFF', end_type='num', end_value=100, end_color='63BE7B'))
        r += 1
    return r
def bxb_matrix(path):
    """Summed bxb counts -> rows ref N, S, V, F, (none) x cols alg n, s, v, (missed)."""
    cols = None; tot = np.zeros(19, np.int64); n = 0
    for line in open(path):
        t = line.split()
        if t and t[0] == 'Record': cols = [c.rstrip("'") for c in t[1:20]]; continue
        if cols and len(t) >= 20 and t[0] not in ('Gross', 'Average', 'Total', 'Sum'):
            try: tot += np.array([int(x) for x in t[1:20]]); n += 1
            except ValueError: pass
    g = dict(zip(cols, tot))
    m = [[g['Nn'], g['Ns'], g['Nv'], g['No']], [g['Sn'], g['Ss'], g['Sv'], g['So']], [g['Vn'], g['Vs'], g['Vv'], g['Vo']],
         [g['Fn'], g['Fs'], g['Fv'], g['Fo']], [g['On'], g['Os'], g['Ov'], 0]]
    return m, n
BROWS = ['N', 'S', 'V', 'F (fusion)', '- (phát hiện thừa)']; BCOLS = ['N', 'S', 'V', 'Bỏ sót']
wb = Workbook()
# ============ EC57 summary ============
ws = wb.active; ws.title = 'EC57 Summary'
r = title(ws, 1, f'EC57 - checkpoint {key}', 'epicmp (rhythm, Gross) và bxb (beat, Gross) - decode chuẩn, hop 5 s + TTA. Cột xám = model gốc 1m e26 để đối chiếu.' if bk else 'epicmp (rhythm) và bxb (beat) - decode chuẩn, hop 5 s + TTA')
ws.cell(r, 1, 'A. Rhythm - epicmp').font = Font(name='Arial', bold=True); r += 1
rh = {(x['db'], x['class']): x for x in rd(f'{S}/{key}/rhythm_ec57_summary.csv')}
rb = {(x['db'], x['class']): x for x in rd(f'{S}/{bk}/rhythm_ec57_summary.csv')} if bk else {}
cols = ['D_Se', 'D_+P', 'D_F1', 'E_Se', 'E_+P', 'E_F1']
head(ws, r, ['Database', 'Lớp', 'Record', 'D Se', 'D +P', 'D F1', 'E Se', 'E +P', 'E F1'] + (['Gốc D F1', 'Gốc E F1'] if bk else [])); r += 1
for db in DBS:
    for c in ['AFIB', 'SVT', 'VT', 'AVB']:
        x = rh.get((db, c))
        if not x: continue
        put(ws, r, 1, db, '@', True); put(ws, r, 2, c, '@', True); put(ws, r, 3, int(x['records']), '0')
        for j, k in enumerate(cols): put(ws, r, 4 + j, num(x[k]))
        if bk and (db, c) in rb:
            put(ws, r, 10, num(rb[(db, c)]['D_F1']), font=GR); put(ws, r, 11, num(rb[(db, c)]['E_F1']), font=GR)
        r += 1
r += 1; ws.cell(r, 1, 'B. Beat - bxb').font = Font(name='Arial', bold=True); r += 1
bh = {x['db']: x for x in rd(f'{S}/{key}/beats/ec57_summary.csv')}
bb = {x['db']: x for x in rd(f'{S}/{bk}/beats/ec57_summary.csv')} if bk else {}
bc = ['Q_Se', 'Q_+P', 'V_Se', 'V_+P', 'S_Se', 'S_+P']
head(ws, r, ['Database', 'Record', 'QRS Se', 'QRS +P', 'VEB Se', 'VEB +P', 'SVEB Se', 'SVEB +P', 'Tổng QRS'] + (['Gốc VEB Se', 'Gốc VEB +P', 'Gốc SVEB Se', 'Gốc SVEB +P'] if bk else [])); r += 1
for db in DBS:
    x = bh.get(db)
    if not x: continue
    put(ws, r, 1, db, '@', True); put(ws, r, 2, int(x['records']), '0')
    for j, k in enumerate(bc): put(ws, r, 3 + j, num(x[k]))
    put(ws, r, 9, int(x['total_QRS']), '#,##0')
    if bk and db in bb:
        for j, k in enumerate(['V_Se', 'V_+P', 'S_Se', 'S_+P']): put(ws, r, 10 + j, num(bb[db][k]), font=GR)
    r += 1
ws.cell(r + 1, 1, 'afdb: tham chiếu beat = .qrs (chỉ vị trí QRS, không phân loại). "-" = database không có beat V/S.').font = Font(name='Arial', italic=True, color='595959')
widths(ws, 12, 10, 16)
# ============ EC57 confusion ============
cmj = json.load(open(f'{S}/{key}/ec57_confusion.json'))
ws = wb.create_sheet('EC57 Rhythm confusion')
r = title(ws, 1, 'EC57 - rhythm confusion theo giây (sau decode chuẩn)', 'Khoảng chấm giống epicmp: từ phút thứ 5 đến hết record. Tham chiếu = rhythm mark .atr ((AFL gộp vào AFIB, mã khác = SINUS).')
for db in DBS:
    R = np.array(cmj[db]['rhythm_seconds'], float); keep = [i for i in range(5) if R[i].sum() > 0 or R[:, i].sum() > 0]
    r = matrix(ws, r, f'{db} ({cmj[db]["records"]} record)', R[np.ix_(keep, keep)], [C[i] for i in keep], [C[i] for i in keep], lambda i, j: i == j, 'Tổng (s)')
widths(ws, 30)
ws = wb.create_sheet('EC57 Beat confusion')
r = title(ws, 1, 'EC57 - beat confusion = ma trận đếm của bxb (Nn, Ns, Nv, ...) cộng qua các record', 'Đúng số liệu bxb: từ phút thứ 5, khớp 150 ms, quy tắc AAMI. Hàng F = beat fusion tham chiếu; hàng "-" = phát hiện thừa; cột "Bỏ sót" = không được phát hiện. Se/P+ ở đây là theo ma trận (bxb tính +P theo AAMI, có thể lệch nhẹ ở VEB do beat F).')
for db in DBS:
    f = f'{S}/{key}/beats/{db}/{db}_QRS_report_line.out'
    try: m, n = bxb_matrix(f)
    except OSError: continue
    r = matrix(ws, r, f'{db} ({n} record, tham chiếu .{cmj[db]["beat_ref_ext"]})', m, BROWS, BCOLS, lambda i, j: i == j and i < 3)
widths(ws, 40)
# ============ dataset-eval v4 ============
v = json.load(open(f'{S}/{key}/dataset_eval_v4/result.json')); vb = json.load(open(f'{S}/{bk}/dataset_eval_v4/result.json')) if bk else None
ws = wb.create_sheet('v4 Summary')
r = title(ws, 1, f'dataset-eval/v4 - checkpoint {key}', 'beat-eval-dataset 5,227 strip (tham chiếu .atr); rhythm-eval-dataset 4,678 event (trong vùng đánh dấu). Cột xám = model gốc.')
ws.cell(r, 1, 'A. Beat - bxb').font = Font(name='Arial', bold=True); r += 1
head(ws, r, ['Phạm vi', 'QRS Se', 'QRS +P', 'VEB Se', 'VEB +P', 'SVEB Se', 'SVEB +P'] + (['Gốc VEB Se', 'Gốc VEB +P', 'Gốc SVEB Se', 'Gốc SVEB +P'] if bk else [])); r += 1
for prot, pl in (('mark_window', 'Vùng đánh dấu (chuẩn beat team)'), ('whole_strip', 'Toàn strip 60 s')):
    x = v['beats']['bxb'][prot]; put(ws, r, 1, pl, '@', True)
    for j, k in enumerate(bc): put(ws, r, 2 + j, num(x.get(k)))
    if bk:
        for j, k in enumerate(['V_Se', 'V_+P', 'S_Se', 'S_+P']): put(ws, r, 8 + j, num(vb['beats']['bxb'][prot].get(k)), font=GR)
    r += 1
r += 1; ws.cell(r, 1, 'B. Rhythm - epicmp (Gross, trong vùng đánh dấu)').font = Font(name='Arial', bold=True); r += 1
TGT = {'AFIB': (90, 94, 88, 96), 'SVT': (89, 88, 60, 78), 'VT': (81, 90, 62, 86), 'AVB': (25, 77, 31, 65)}
head(ws, r, ['Lớp', 'E Se', 'E +P', 'D Se', 'D +P', '5.4.1 E Se', '5.4.1 E +P', '5.4.1 D Se', '5.4.1 D +P'] + (['Gốc E Se', 'Gốc E +P', 'Gốc D Se', 'Gốc D +P'] if bk else [])); r += 1
for c in ['AFIB', 'SVT', 'VT', 'AVB']:
    put(ws, r, 1, c, '@', True)
    for j, x in enumerate(v['rhythm']['epicmp'][c]): put(ws, r, 2 + j, x)
    for j, x in enumerate(TGT[c]): put(ws, r, 6 + j, x, f=TF)
    if bk:
        for j, x in enumerate(vb['rhythm']['epicmp'][c]): put(ws, r, 10 + j, x, font=GR)
    r += 1
r += 1; ws.cell(r, 1, 'C. Rhythm - theo event (Se / PPV / F1)').font = Font(name='Arial', bold=True); r += 1
head(ws, r, ['Lớp', 'Se', 'PPV', 'F1'] + (['Gốc Se', 'Gốc PPV', 'Gốc F1'] if bk else [])); r += 1
for c in ['AFIB', 'SVT', 'VT', 'AVB']:
    put(ws, r, 1, c, '@', True)
    for j, x in enumerate(v['rhythm']['event_stats'][c]): put(ws, r, 2 + j, x)
    if bk:
        for j, x in enumerate(vb['rhythm']['event_stats'][c]): put(ws, r, 5 + j, x, font=GR)
    r += 1
widths(ws, 32, 11, 15)
ws = wb.create_sheet('v4 Beat confusion')
r = title(ws, 1, 'dataset-eval/v4 beat - confusion = ma trận đếm của bxb (beat sau post-process)', 'Hàng F = beat fusion tham chiếu; hàng "-" = phát hiện thừa; cột "Bỏ sót" = không được phát hiện.')
for prot, pl in (('mark_window', 'Vùng đánh dấu (chuẩn beat team)'), ('whole_strip', 'Toàn strip 60 s')):
    m, n = bxb_matrix(f'{S}/{key}/dataset_eval_v4/beats/bxb_{prot}/beat_eval_v4/beat_eval_v4_QRS_report_line.out')
    r = matrix(ws, r, f'beat-eval-dataset - {pl} ({n} record)', m, BROWS, BCOLS, lambda i, j: i == j and i < 3)
widths(ws, 34)
ws = wb.create_sheet('v4 Rhythm confusion')
r = title(ws, 1, 'dataset-eval/v4 rhythm - confusion', 'Theo sample: lưới 40 ms trong vùng đánh dấu, sau decode chuẩn (AVB2/AVB3 đúng khi dự đoán AVB). Theo event: 1 event = 1 đơn vị.')
PN = ['SINUS', 'AFIB', 'SVT', 'VT', 'AVB', 'NOISE']; RT = ['SINUS', 'AFIB', 'SVT', 'VT', 'AVB2', 'AVB3']
sc = np.array(v['rhythm']['sample_cm']); keep = [i for i, t in enumerate(RT) if sc[i].sum() > 0]
r = matrix(ws, r, 'Theo sample (số bước 40 ms)', sc[keep], [RT[i] for i in keep], PN, lambda i, j: PN[j] == ('AVB' if RT[keep[i]] in ('AVB2', 'AVB3') else RT[keep[i]]))
EC = ['AFIB', 'SVT', 'VT', 'AVB']; ER = ['AFIB', 'SVT', 'VT', 'AVB2', 'AVB3']
r = matrix(ws, r, 'Theo event', [[v['rhythm']['event_cm'][t][c] for c in EC] for t in ER], ER, EC, lambda i, j: EC[j] == ['AFIB', 'SVT', 'VT', 'AVB', 'AVB'][i], 'Tổng event')
widths(ws, 30)
out = f'{S}/{key}/{key}_report.xlsx'; wb.save(out); print('saved', out, wb.sheetnames)
